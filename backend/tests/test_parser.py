import pytest
import io
import docx
import pymupdf
from app.services.parser import (
    parse_resume_document,
    extract_contact_info,
    analyze_achievements_and_metrics,
    check_formatting_and_parsing
)

def test_extract_contact_info():
    sample_text = """
    Alex Johnson
    alex.johnson@example.com | (555) 123-4567 | San Francisco, CA
    linkedin.com/in/alexjohnson | github.com/alexjohnson
    """
    contact = extract_contact_info(sample_text)
    assert contact["email"] == "alex.johnson@example.com"
    assert "123-4567" in contact["phone"]
    assert "linkedin.com/in/alexjohnson" in contact["linkedin"]
    assert "github.com/alexjohnson" in contact["github"]

def test_achievements_and_metrics():
    sample_bullet = """
    - Spearheaded migration of payment services, reducing latency by 45% and saving $120,000 annually.
    - Led team of 6 engineers across 3 production microservices.
    """
    achieve = analyze_achievements_and_metrics(sample_bullet)
    assert achieve["metric_count"] >= 2
    assert achieve["unique_action_verbs_count"] >= 2
    assert "spearheaded" in achieve["action_verbs"] or "led" in achieve["action_verbs"]

def test_parse_txt_document():
    txt_content = b"""
    Jane Doe
    jane.doe@example.com | (555) 987-6543
    
    Professional Summary
    Experienced Full-Stack Engineer with 6 years building distributed cloud applications.
    
    Technical Skills
    Python, React, TypeScript, PostgreSQL, Docker, AWS
    
    Professional Experience
    Senior Software Engineer - Acme Corp (2021 - Present)
    - Architected real-time analytics pipeline processing 50M events per day.
    - Reduced API response time by 35% using Redis caching.
    
    Education
    B.S. in Computer Science - University of California
    """
    res = parse_resume_document(txt_content, "resume.txt")
    assert res["contact"]["email"] == "jane.doe@example.com"
    assert "Python" in res["raw_text"]
    assert res["word_count"] > 30

def test_parse_docx_document():
    doc = docx.Document()
    doc.add_heading("John Smith", level=1)
    doc.add_paragraph("john.smith@domain.com | (555) 234-5678")
    doc.add_paragraph("Experience: Built scalable APIs using FastAPI and Python.")
    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    docx_bytes = doc_stream.getvalue()
    
    res = parse_resume_document(docx_bytes, "resume.docx")
    assert res["contact"]["email"] == "john.smith@domain.com"
    assert "FastAPI" in res["raw_text"]

def test_parse_pdf_document():
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Sarah Connor\nsarah@cyberdyne.org | 555-0199\nSkills: Python, Kubernetes, Docker")
    pdf_bytes = doc.tobytes()
    doc.close()
    
    res = parse_resume_document(pdf_bytes, "sarah_resume.pdf")
    assert res["contact"]["email"] == "sarah@cyberdyne.org"
    assert "Kubernetes" in res["raw_text"]

def test_empty_document_raises_error():
    with pytest.raises(ValueError):
        parse_resume_document(b"   \n\t  ", "empty.txt")
